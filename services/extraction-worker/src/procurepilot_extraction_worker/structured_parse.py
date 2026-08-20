from __future__ import annotations

import csv
import io
from collections.abc import Iterable

from procurepilot_extraction_worker.models import ExtractedField, ExtractedLine, ExtractionResult
from procurepilot_extraction_worker.validation import decimal_from_extracted

STRUCTURED_MODEL_VERSION = "structured-quotation-parser-v1"
CSV_MIME_TYPES = {"text/csv", "application/vnd.ms-excel"}
XLSX_MIME_TYPES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def is_structured_mime_type(mime_type: str) -> bool:
    return mime_type in CSV_MIME_TYPES | XLSX_MIME_TYPES


def parse_structured_content(content: bytes, *, mime_type: str) -> ExtractionResult:
    if mime_type in XLSX_MIME_TYPES:
        return _parse_xlsx(content)
    text = content.decode("utf-8-sig")
    return _parse_rows(csv.DictReader(io.StringIO(text)))


def _parse_xlsx(content: bytes) -> ExtractionResult:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return _parse_rows([])
    headers = [str(value or "").strip() for value in rows[0]]
    dict_rows = [
        {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
        for row in rows[1:]
    ]
    return _parse_rows(dict_rows)


def _parse_rows(rows: Iterable[dict[str, object]]) -> ExtractionResult:
    header: dict[str, ExtractedField] = {}
    lines: list[ExtractedLine] = []
    stated_total: ExtractedField | None = None
    for index, row in enumerate(rows, start=1):
        normalised = {str(key).strip().lower(): value for key, value in row.items()}
        currency = str(normalised.get("currency") or "GBP").strip().upper()
        if "currency" not in header:
            header["currency"] = _field("currency", currency)
        for name in ("supplier_name", "issue_date", "expiry_date"):
            value = normalised.get(name)
            if value not in {None, ""} and name not in header:
                header[name] = _field(name, str(value))
        if normalised.get("stated_total") not in {None, ""}:
            stated_total = _field(
                "stated_total",
                {"amount": _decimal_string(normalised["stated_total"]), "currency": currency},
            )
        description = str(normalised.get("description") or normalised.get("original_text") or "")
        fields: dict[str, ExtractedField] = {}
        for name in ("quantity", "pack_count", "unit_size", "pack_unit", "vat_rate"):
            value = normalised.get(name)
            if value not in {None, ""}:
                fields[name] = _field(name, _normalise_scalar(name, value))
        for name in ("unit_price", "delivery_fee", "discount"):
            value = normalised.get(name)
            if value not in {None, ""}:
                fields[name] = _field(
                    name, {"amount": _decimal_string(value), "currency": currency}
                )
        lines.append(
            ExtractedLine(
                line_number=int(normalised.get("line_number") or index),
                original_text=description or f"Structured row {index}",
                fields=fields,
            )
        )
    return ExtractionResult(
        method="structured_parse",
        model_version=STRUCTURED_MODEL_VERSION,
        header=header,
        lines=lines,
        stated_total=stated_total,
        cost_usd=0.0,
    )


def _field(name: str, value: object) -> ExtractedField:
    return ExtractedField(
        name=name,
        value=value,
        confidence=1.0,
        source_page=None,
        source_region=None,
    )


def _normalise_scalar(name: str, value: object) -> object:
    if name == "pack_count":
        return int(decimal_from_extracted(value))
    if name in {"quantity", "unit_size", "vat_rate"}:
        return _decimal_string(value)
    return str(value).strip()


def _decimal_string(value: object) -> str:
    return format(decimal_from_extracted(value), "f")
